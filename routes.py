import os
import io
import csv
import json
import zipfile
import pandas as pd
from datetime import datetime
from flask import (Blueprint, render_template, redirect, url_for,
                   flash, request, make_response, current_app, send_file)
from flask_login import login_user, login_required, logout_user, current_user
from werkzeug.utils import secure_filename

# Import dari file model dan utility kamu
from models import db, User, Audit
from utils import process_image
from fuzzy_logic import calculate_skor_fuzzy

massal_progress = {}

main = Blueprint('main', __name__)

# ─── Helper ──────────────────────────────────────────────────────────────────

def _audit_to_result(audit):
    """Konversi objek Audit dari DB ke dict result yang dipakai template."""
    return {
        'id':     audit.id,
        'lokasi': audit.lokasi,
        'lat':    audit.lat,
        'long':   audit.long,
        'waktu':  audit.tanggal.strftime("%d-%m-%Y %H:%M"),
        'skor':   audit.skor,
        'status': audit.status,
        'images': {
            'depan': audit.img_depan or '',
            'kanan': audit.img_kanan or '',
            'kiri':  audit.img_kiri  or '',
        },
        'counts': {
            'depan': {
                'orang': audit.d_orang or 0,
                'bangkutaman': audit.d_bangku or 0,
                'tempatsampah': audit.d_sampah or 0,
                'pohonpeneduh': audit.d_pohon or 0,
                'lampulalulintas': audit.d_lampu or 0,
            },
            'kanan': {
                'orang': audit.k_orang or 0,
                'bangkutaman': audit.k_bangku or 0,
                'tempatsampah': audit.k_sampah or 0,
                'pohonpeneduh': audit.k_pohon or 0,
                'lampulalulintas': audit.k_lampu or 0,
            },
            'kiri': {
                'orang': audit.ki_orang or 0,
                'bangkutaman': audit.ki_bangku or 0,
                'tempatsampah': audit.ki_sampah or 0,
                'pohonpeneduh': audit.ki_pohon or 0,
                'lampulalulintas': audit.ki_lampu or 0,
            },
        },
    }

def normalize_coord(c_str, is_lat=True):
    """Membagi angka ribuan/jutaan dengan 10 hingga menjadi format koordinat GPS standar."""
    if not c_str: return ''
    try:
        c_str = str(c_str).replace(',', '.').strip()
        val = float(c_str)
        limit = 90.0 if is_lat else 180.0
        while abs(val) > limit and abs(val) > 0:
            val = val / 10.0
        return str(round(val, 7))
    except ValueError:
        return str(c_str)


# ─── Auth ─────────────────────────────────────────────────────────────────────

@main.route('/')
def index():
    return redirect(url_for('main.login'))


@main.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('main.dashboard'))
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user)
            return redirect(url_for('main.dashboard'))
        else:
            flash('Login Gagal. Periksa username dan password.', 'danger')
    return render_template('login.html')


@main.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('main.login'))


# ─── Dashboard ────────────────────────────────────────────────────────────────

@main.route('/dashboard')
@login_required
def dashboard():
    audits = Audit.query.filter_by(user_id=current_user.id).order_by(Audit.tanggal.asc()).all()

    raw_data = [{
        'hari': a.tanggal.strftime("%d %b"),
        'bulan': a.tanggal.strftime("%b %Y"),
        'tanggal_asli': a.tanggal.strftime("%Y-%m-%d %H:%M:%S"),
        'skor': a.skor
    } for a in audits]

    if not raw_data:
        raw_data = [{'hari': 'Belum ada data', 'bulan': 'Belum ada data', 'tanggal_asli': '', 'skor': 0}]

    total_audit   = Audit.query.filter_by(user_id=current_user.id).count()
    lokasi_layak  = Audit.query.filter(Audit.user_id == current_user.id, Audit.status.in_(['Pemeliharaan', 'Perencanaan'])).count()
    lokasi_kurang = Audit.query.filter(Audit.user_id == current_user.id, Audit.status.in_(['Monitoring', 'Intervensi'])).count()

    return render_template('dashboard.html',
                           name=current_user.username,
                           raw_data=raw_data,
                           total_audit=total_audit,
                           lokasi_layak=lokasi_layak,
                           lokasi_kurang=lokasi_kurang)


# ─── Analysis ─────────────────────────────────────────────────────────────────

@main.route('/standar')
@login_required
def standar():
    """Tampilkan referensi acuan standar audit."""
    return render_template('acuan_standar.html')

@main.route('/analysis', methods=['GET', 'POST'])
@login_required
def analysis():
    result = None
    upload_folder = current_app.config['UPLOAD_FOLDER']

    if request.method == 'POST':
        jenis_upload = request.form.get('jenis_upload')

        # --- MODE SATUAN (Urutan: Depan, Kanan, Kiri) ---
        if jenis_upload == 'satuan':
            lokasi = request.form.get('lokasi')
            lat    = request.form.get('lat')
            long   = request.form.get('long')

            f_depan, c_depan = process_image(request.files.get('img_depan'), 'depan', upload_folder)
            f_kanan, c_kanan = process_image(request.files.get('img_kanan'), 'kanan', upload_folder)
            f_kiri,  c_kiri  = process_image(request.files.get('img_kiri'),  'kiri',  upload_folder)

            # Hitung jumlah gambar yang valid diproses
            image_count = 0
            if f_depan is not None: image_count += 1
            if f_kanan is not None: image_count += 1
            if f_kiri is not None: image_count += 1

            if image_count == 0:
                flash('Minimal harus ada 1 foto yang berhasil diproses.', 'danger')
                return redirect(url_for('main.analysis'))

            skor_fuzzy, status_fuzzy = calculate_skor_fuzzy(c_depan, c_kanan, c_kiri, image_count)

            lat_norm = normalize_coord(lat, is_lat=True)
            long_norm = normalize_coord(long, is_lat=False)

            audit_baru = Audit(  # type: ignore[call-arg]
                lokasi=lokasi, lat=lat_norm, long=long_norm,
                skor=float(skor_fuzzy), status=status_fuzzy,
                author=current_user,
                # Simpan nama file gambar (atau string kosong jika tidak ada)
                img_depan=f_depan or '', img_kanan=f_kanan or '', img_kiri=f_kiri or '',
                # Simpan counts Depan
                d_orang=c_depan.get('orang', 0), d_bangku=c_depan.get('bangkutaman', 0),
                d_sampah=c_depan.get('tempatsampah', 0), d_pohon=c_depan.get('pohonpeneduh', 0),
                d_lampu=c_depan.get('lampulalulintas', 0),
                # Simpan counts Kanan
                k_orang=c_kanan.get('orang', 0), k_bangku=c_kanan.get('bangkutaman', 0),
                k_sampah=c_kanan.get('tempatsampah', 0), k_pohon=c_kanan.get('pohonpeneduh', 0),
                k_lampu=c_kanan.get('lampulalulintas', 0),
                # Simpan counts Kiri
                ki_orang=c_kiri.get('orang', 0), ki_bangku=c_kiri.get('bangkutaman', 0),
                ki_sampah=c_kiri.get('tempatsampah', 0), ki_pohon=c_kiri.get('pohonpeneduh', 0),
                ki_lampu=c_kiri.get('lampulalulintas', 0),
            )
            db.session.add(audit_baru)
            db.session.commit()

            result = _audit_to_result(audit_baru)
            flash('Audit Lokasi berhasil diproses!', 'success')

        elif jenis_upload == 'massal':
            user_id_str = str(current_user.id)
            try:
                file_csv = request.files.get('file_csv')
                file_zip = request.files.get('file_zip')
                if file_csv and file_zip:
                    zip_filename = file_zip.filename or 'upload.zip'
                    zip_path    = os.path.join(upload_folder, secure_filename(zip_filename))
                    file_zip.save(zip_path)
                    extract_dir = os.path.join(upload_folder, 'extracted_' + datetime.now().strftime("%H%M%S"))
                    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                        zip_ref.extractall(extract_dir)

                    # BUG FIX #3: Jika hasil ekstraksi mengandung subfolder, arahkan ke subfolder pertama
                    subfolders = [f.path for f in os.scandir(extract_dir) if f.is_dir()]
                    if subfolders:
                        extract_dir = subfolders[0]

                    # Baca Excel (.xlsx/.xls) atau CSV via pandas — gunakan BytesIO/StringIO agar aman
                    raw_bytes = file_csv.stream.read()
                    csv_filename = file_csv.filename or ''
                    if csv_filename.lower().endswith(('.xlsx', '.xls')):
                        df = pd.read_excel(io.BytesIO(raw_bytes))
                    else:
                        try:
                            decoded_csv = raw_bytes.decode('utf-8-sig')
                        except UnicodeDecodeError:
                            decoded_csv = raw_bytes.decode('latin-1')
                        df = pd.read_csv(io.StringIO(decoded_csv))
                    # Normalisasi nama kolom: strip & lowercase
                    df.columns = [c.strip().lower() for c in df.columns]
                    rows = df.to_dict('records')

                    total_rows = len(rows)
                    
                    massal_progress[user_id_str] = {'status': 'Mengekstrak file ZIP...', 'progress': 5}

                    berhasil = 0
                    gagal = 0

                    VALID_EXTS = ['.jpg', '.jpeg', '.png', '.bmp', '.webp', '.jfif']

                    def cari_file_by_lokasi(base_path, lokasi, arah):
                        """Merakit nama file dinamis dari lokasi+arah (misal: titik001_depan.jpg)
                        dan mencari secara rekursif ke seluruh sub-folder."""
                        for ext in VALID_EXTS:
                            nama_target = f"{lokasi}_{arah}{ext}"
                            for root, dirs, files in os.walk(base_path):
                                for file in files:
                                    if file.lower().strip() == nama_target:
                                        return os.path.join(root, file)
                        return None

                    for index, raw_row in enumerate(rows):
                        try:
                            # Bersihkan setiap nilai cell; tangani NaN dari pandas
                            row = {k: str(v).strip() if pd.notna(v) else '' for k, v in raw_row.items()}

                            # Kolom lokasi: cek 'location' (Excel baru) ATAU 'lokasi' (format lama)
                            lokasi_val = row.get('location') or row.get('lokasi') or row.get('nama lokasi') or f"Titik_{index+1}"
                            lokasi_clean = lokasi_val.strip().lower()  # misal: "titik001"

                            massal_progress[user_id_str] = {'status': f'Menganalisis foto dengan YOLOv11... ({index+1}/{total_rows})', 'progress': 10 + int(80 * (index / total_rows))}

                            # Cari gambar dengan merakit nama dinamis: {lokasi_clean}_{arah}.{ext}
                            f_depan_path = cari_file_by_lokasi(extract_dir, lokasi_clean, 'depan')
                            f_kanan_path = cari_file_by_lokasi(extract_dir, lokasi_clean, 'kanan')
                            f_kiri_path  = cari_file_by_lokasi(extract_dir, lokasi_clean, 'kiri')

                            # Panggil process_image dengan None atau string dari f_*_path
                            f_depan, c_depan = process_image(f_depan_path, 'depan', upload_folder, is_path=True) if f_depan_path else (None, {'orang': 0, 'bangkutaman': 0, 'tempatsampah': 0, 'pohonpeneduh': 0, 'lampulalulintas': 0})
                            f_kanan, c_kanan = process_image(f_kanan_path, 'kanan', upload_folder, is_path=True) if f_kanan_path else (None, {'orang': 0, 'bangkutaman': 0, 'tempatsampah': 0, 'pohonpeneduh': 0, 'lampulalulintas': 0})
                            f_kiri,  c_kiri  = process_image(f_kiri_path, 'kiri',  upload_folder, is_path=True) if f_kiri_path else (None, {'orang': 0, 'bangkutaman': 0, 'tempatsampah': 0, 'pohonpeneduh': 0, 'lampulalulintas': 0})

                            image_count = 0
                            if f_depan is not None: image_count += 1
                            if f_kanan is not None: image_count += 1
                            if f_kiri is not None: image_count += 1

                            if image_count == 0:
                                gagal += 1
                                print(f"Gagal proses baris {index+1}: Semua arah gambar tidak ditemukan untuk {lokasi_val}")
                                continue  # Skip baris jika SEMUA gambar gagal/hilang

                            massal_progress[user_id_str] = {'status': f'Menghitung skor Fuzzy Mamdani... ({index+1}/{total_rows})', 'progress': 10 + int(80 * ((index + 0.5) / total_rows))}
                            skor_fuzzy, status_fuzzy = calculate_skor_fuzzy(c_depan, c_kanan, c_kiri, image_count)

                            # Normalisasi koordinat
                            lat_norm = normalize_coord(row.get('latitude') or row.get('lat'), is_lat=True)
                            long_norm = normalize_coord(row.get('longitude') or row.get('long'), is_lat=False)

                            massal_progress[user_id_str] = {'status': f'Menyimpan hasil ke database... ({index+1}/{total_rows})', 'progress': 10 + int(80 * ((index + 0.8) / total_rows))}
                            audit_baru = Audit(  # type: ignore[call-arg]
                                lokasi=lokasi_val, 
                                lat=lat_norm, 
                                long=long_norm,
                                skor=float(skor_fuzzy), status=status_fuzzy, author=current_user,
                                img_depan=f_depan or '', img_kanan=f_kanan or '', img_kiri=f_kiri or '',
                                d_orang=c_depan.get('orang', 0), d_bangku=c_depan.get('bangkutaman', 0),
                                d_sampah=c_depan.get('tempatsampah', 0), d_pohon=c_depan.get('pohonpeneduh', 0),
                                d_lampu=c_depan.get('lampulalulintas', 0),
                                k_orang=c_kanan.get('orang', 0), k_bangku=c_kanan.get('bangkutaman', 0),
                                k_sampah=c_kanan.get('tempatsampah', 0), k_pohon=c_kanan.get('pohonpeneduh', 0),
                                k_lampu=c_kanan.get('lampulalulintas', 0),
                                ki_orang=c_kiri.get('orang', 0), ki_bangku=c_kiri.get('bangkutaman', 0),
                                ki_sampah=c_kiri.get('tempatsampah', 0), ki_pohon=c_kiri.get('pohonpeneduh', 0),
                                ki_lampu=c_kiri.get('lampulalulintas', 0),
                            )
                            db.session.add(audit_baru)
                            db.session.commit()
                            berhasil += 1
                        except Exception as inner_e:
                            db.session.rollback()
                            gagal += 1
                            print(f"Error baris {index+1}: {inner_e}")

                    massal_progress[user_id_str] = {'status': 'Menyelesaikan...', 'progress': 100}

                    if gagal > 0:
                        flash(f'Audit Massal Selesai: {berhasil} berhasil, {gagal} gagal. (Periksa log untuk detail)', 'warning')
                    else:
                        flash(f'Audit Massal Selesai: {berhasil} data berhasil diproses!', 'success')
                    return redirect(url_for('main.history'))
            except Exception as e:
                import traceback
                error_trace = traceback.format_exc()
                print(error_trace)
                massal_progress[user_id_str] = {'status': 'Error', 'progress': 100}
                flash(f'Terjadi Kesalahan Sistem (Internal Error) saat memproses Mass Audit: {str(e)}', 'danger')
                return redirect(url_for('main.analysis'))

    return render_template('analysis.html', result=result)

@main.route('/massal_progress')
@login_required
def get_massal_progress():
    user_id_str = str(current_user.id)
    return current_app.response_class(
        json.dumps(massal_progress.get(user_id_str, {'status': '', 'progress': 0})),
        mimetype='application/json'
    )


@main.route('/view/<int:id>')
@login_required
def view_audit(id):
    """Tampilkan ulang hasil analisis YOLO dari riwayat."""
    audit = Audit.query.get_or_404(id)
    if audit.author != current_user:
        flash('Akses ditolak.', 'danger')
        return redirect(url_for('main.history'))
    result = _audit_to_result(audit)
    return render_template('analysis.html', result=result)


# ─── History ──────────────────────────────────────────────────────────────────

@main.route('/history')
@login_required
def history():
    riwayat_data = Audit.query.filter_by(user_id=current_user.id).order_by(Audit.tanggal.desc()).all()
    all_audits   = Audit.query.filter_by(user_id=current_user.id).all()
    return render_template('history.html', audits=riwayat_data, all_audits=all_audits)


@main.route('/delete/<int:id>')
@login_required
def delete_audit(id):
    audit_item = Audit.query.get_or_404(id)
    if audit_item.author == current_user:
        db.session.delete(audit_item)
        db.session.commit()
        flash('Data berhasil dihapus.', 'success')
    return redirect(url_for('main.history'))


@main.route('/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_audit(id):
    audit_item = Audit.query.get_or_404(id)
    if audit_item.author != current_user:
        flash('Anda tidak memiliki akses ke data ini.', 'danger')
        return redirect(url_for('main.history'))
    if request.method == 'POST':
        audit_item.lokasi = request.form.get('lokasi', audit_item.lokasi)
        audit_item.lat    = request.form.get('lat',    audit_item.lat)
        audit_item.long   = request.form.get('long',   audit_item.long)
        db.session.commit()
        flash('Data audit berhasil diperbarui.', 'success')
        return redirect(url_for('main.history'))
    return render_template('edit_audit.html', audit=audit_item)


# ─── Download Reports ─────────────────────────────────────────────────────────

def get_filtered_audits(user_id):
    search = request.args.get('search', '').lower()
    status_filter = request.args.get('status', '')

    query = Audit.query.filter_by(user_id=user_id)
    if status_filter:
        query = query.filter_by(status=status_filter)
    
    audits = query.order_by(Audit.tanggal.desc()).all()
    if search:
        audits = [a for a in audits if search in (a.lokasi or '').lower()]
    return audits, search, status_filter

STATUS_EMOJI = {'Pemeliharaan': '✅', 'Monitoring': '⚠️', 'Perencanaan': '📅', 'Intervensi': '🚨'}
STATUS_COLORS_HEX = {'Pemeliharaan': '10B981', 'Monitoring': 'F59E0B', 'Perencanaan': '64748B', 'Intervensi': 'EF4444'}
STATUS_COLORS_RGB = {
    'Pemeliharaan': (16, 185, 129),
    'Monitoring': (245, 158, 11),
    'Perencanaan': (100, 116, 139),
    'Intervensi': (239, 68, 68)
}

@main.route('/download/csv')
@login_required
def download_report():
    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow(['No', 'Tanggal & Waktu', 'Nama Lokasi', 'Koordinat', 'Objek Terdeteksi', 'Skor', 'Status'])
    audits, _, _ = get_filtered_audits(current_user.id)
    for i, audit in enumerate(audits, 1):
        status_text = f"{STATUS_EMOJI.get(audit.status, '')} {audit.status}"
        koordinat = f"{audit.lat}, {audit.long}"
        orang = (audit.d_orang or 0) + (audit.k_orang or 0) + (audit.ki_orang or 0)
        bangku = (audit.d_bangku or 0) + (audit.k_bangku or 0) + (audit.ki_bangku or 0)
        sampah = (audit.d_sampah or 0) + (audit.k_sampah or 0) + (audit.ki_sampah or 0)
        pohon = (audit.d_pohon or 0) + (audit.k_pohon or 0) + (audit.ki_pohon or 0)
        lampu = (audit.d_lampu or 0) + (audit.k_lampu or 0) + (audit.ki_lampu or 0)
        objek = f"🚶 {orang}   🪑 {bangku}   🗑️ {sampah}   🌳 {pohon}   🚦 {lampu}"
        cw.writerow([i, audit.tanggal.strftime('%d/%m/%Y %H:%M'), audit.lokasi,
                     koordinat, objek, audit.skor, status_text])
    output = make_response(si.getvalue().encode('utf-8-sig')) # Use utf-8-sig to support emojis in excel
    output.headers["Content-Disposition"] = "attachment; filename=Laporan_Audit.csv"
    output.headers["Content-type"] = "text/csv; charset=utf-8"
    return output


@main.route('/download/excel')
@login_required
def download_excel():
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except ImportError:
        flash('Library openpyxl belum terinstall. Jalankan: pip install openpyxl', 'danger')
        return redirect(url_for('main.history'))

    wb  = openpyxl.Workbook()
    ws  = wb.active
    assert ws is not None  # wb.active is always valid for a new Workbook
    ws.title = "Laporan Audit"

    audits, search, status_filter = get_filtered_audits(current_user.id)
    filter_info = f"Filter Aktif: {status_filter or 'Semua Status'}"
    if search: filter_info += f" | Pencarian: '{search}'"
    
    ws.append([filter_info])
    ws.append([]) # Empty row

    # Header styling
    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(bold=True, color="F1C40F", size=11)
    border      = Border(
        bottom=Side(style='thin', color='444444'),
        right=Side(style='thin', color='333333'),
        left=Side(style='thin', color='333333'),
        top=Side(style='thin', color='333333')
    )
    headers = ['No', 'Tanggal & Waktu', 'Nama Lokasi', 'Koordinat', 'Objek Terdeteksi', 'Skor', 'Status']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border    = border

    for i, audit in enumerate(audits, 1):
        status_text = f"{STATUS_EMOJI.get(audit.status, '')} {audit.status}"
        koordinat = f"{audit.lat}, {audit.long}"
        orang = (audit.d_orang or 0) + (audit.k_orang or 0) + (audit.ki_orang or 0)
        bangku = (audit.d_bangku or 0) + (audit.k_bangku or 0) + (audit.ki_bangku or 0)
        sampah = (audit.d_sampah or 0) + (audit.k_sampah or 0) + (audit.ki_sampah or 0)
        pohon = (audit.d_pohon or 0) + (audit.k_pohon or 0) + (audit.ki_pohon or 0)
        lampu = (audit.d_lampu or 0) + (audit.k_lampu or 0) + (audit.ki_lampu or 0)
        objek = f"🚶 {orang}   🪑 {bangku}   🗑️ {sampah}   🌳 {pohon}   🚦 {lampu}"
        
        row_data = [i, audit.tanggal.strftime('%d/%m/%Y %H:%M'), audit.lokasi,
                    koordinat, objek, audit.skor, status_text]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i + 3, column=col, value=val)
            cell.alignment = Alignment(horizontal='center' if col in [1, 6] else 'left', vertical='center')
            cell.border    = border
            if col == 7: # Status column
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS_HEX.get(audit.status, 'FFFFFF'))
                cell.font = Font(color="FFFFFF", bold=True)

    # Auto-width kolom
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        col_idx = col[0].column
        if isinstance(col_idx, int):
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name='Laporan_Audit.xlsx')


@main.route('/download/word')
@login_required
def download_word():
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import nsdecls
        from docx.oxml import parse_xml

        def set_cell_background(cell, fill_color):
            shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
            cell._tc.get_or_add_tcPr().append(shading_elm)

    except ImportError:
        flash('Library python-docx belum terinstall. Jalankan: pip install python-docx', 'danger')
        return redirect(url_for('main.history'))

    audits, search, status_filter = get_filtered_audits(current_user.id)
    
    doc = Document()

    # Judul
    title_text = 'Laporan Audit Kualitas Ruang Publik'
    if status_filter: title_text += f' - Status: {status_filter}'
    title = doc.add_heading(title_text, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(f'Sistem Audit Cerdas — {current_user.username} — {datetime.now().strftime("%d %B %Y")}')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if search:
        doc.add_paragraph(f"Pencarian: '{search}'").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=7)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    headers   = ['No', 'Tanggal & Waktu', 'Nama Lokasi', 'Koordinat', 'Objek Terdeteksi', 'Skor', 'Status']
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True
        set_cell_background(hdr_cells[i], '1A1A2E')
        run.font.color.rgb = RGBColor(241, 196, 15)

    for i, audit in enumerate(audits, 1):
        row_cells = table.add_row().cells
        status_text = f"{STATUS_EMOJI.get(audit.status, '')} {audit.status}"
        koordinat = f"{audit.lat}, {audit.long}"
        orang = (audit.d_orang or 0) + (audit.k_orang or 0) + (audit.ki_orang or 0)
        bangku = (audit.d_bangku or 0) + (audit.k_bangku or 0) + (audit.ki_bangku or 0)
        sampah = (audit.d_sampah or 0) + (audit.k_sampah or 0) + (audit.ki_sampah or 0)
        pohon = (audit.d_pohon or 0) + (audit.k_pohon or 0) + (audit.ki_pohon or 0)
        lampu = (audit.d_lampu or 0) + (audit.k_lampu or 0) + (audit.ki_lampu or 0)
        objek = f"🚶 {orang}   🪑 {bangku}   🗑️ {sampah}   🌳 {pohon}   🚦 {lampu}"
        
        row_data  = [str(i), audit.tanggal.strftime('%d/%m/%Y %H:%M'), audit.lokasi,
                     koordinat, objek, str(audit.skor), status_text]
        for j, val in enumerate(row_data):
            row_cells[j].text = val
            if j == 6: # Status
                set_cell_background(row_cells[j], STATUS_COLORS_HEX.get(audit.status, 'FFFFFF'))
                # Make text white
                run = row_cells[j].paragraphs[0].runs[0] if row_cells[j].paragraphs[0].runs else row_cells[j].paragraphs[0].add_run(val)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                     as_attachment=True,
                     download_name='Laporan_Audit.docx')


@main.route('/download/pdf')
@login_required
def download_pdf():
    try:
        from fpdf import FPDF  # type: ignore
    except ImportError:
        flash('Library fpdf2 belum terinstall. Jalankan: pip install fpdf2', 'danger')
        return redirect(url_for('main.history'))

    audits, search, status_filter = get_filtered_audits(current_user.id)

    # Some fonts like Helvetica don't support emojis natively in fpdf.
    # We will just append the status text directly for PDF.
    
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Judul
    pdf.set_font('Helvetica', 'B', 16)
    title_text = 'Laporan Audit Kualitas Ruang Publik'
    if status_filter: title_text += f' - Status: {status_filter}'
    pdf.cell(0, 12, title_text, ln=True, align='C')
    pdf.set_font('Helvetica', '', 10)
    pdf.cell(0, 6, f'Sistem Audit Cerdas  |  User: {current_user.username}  |  {datetime.now().strftime("%d %B %Y")}', ln=True, align='C')
    if search:
        pdf.cell(0, 6, f"Pencarian: '{search}'", ln=True, align='C')
    pdf.ln(6)

    # Header tabel
    col_widths = [10, 25, 45, 35, 95, 15, 30] # Removed Aksi, added space to Objek and Status
    headers    = ['No', 'Tanggal', 'Nama Lokasi', 'Koordinat', 'Objek Terdeteksi', 'Skor', 'Status']
    pdf.set_fill_color(26, 26, 46)
    pdf.set_text_color(241, 196, 15)
    pdf.set_font('Helvetica', 'B', 8)
    for w, h in zip(col_widths, headers):
        pdf.cell(w, 8, h, border=1, align='C', fill=True)
    pdf.ln()

    # Baris data
    pdf.set_font('Helvetica', '', 8)
    for i, audit in enumerate(audits, 1):
        koordinat = f"{audit.lat}, {audit.long}"
        orang = (audit.d_orang or 0) + (audit.k_orang or 0) + (audit.ki_orang or 0)
        bangku = (audit.d_bangku or 0) + (audit.k_bangku or 0) + (audit.ki_bangku or 0)
        sampah = (audit.d_sampah or 0) + (audit.k_sampah or 0) + (audit.ki_sampah or 0)
        pohon = (audit.d_pohon or 0) + (audit.k_pohon or 0) + (audit.ki_pohon or 0)
        lampu = (audit.d_lampu or 0) + (audit.k_lampu or 0) + (audit.ki_lampu or 0)
        objek = f"🚶 {orang}   🪑 {bangku}   🗑️ {sampah}   🌳 {pohon}   🚦 {lampu}"
        
        row_data = [str(i), audit.tanggal.strftime('%d/%m/%Y %H:%M'), audit.lokasi,
                    koordinat, objek, str(audit.skor)]
        
        pdf.set_text_color(30, 30, 30)
        # Alternate row color
        if i % 2 == 0:
            pdf.set_fill_color(240, 240, 240)
        else:
            pdf.set_fill_color(255, 255, 255)
            
        for w, val in zip(col_widths[:-1], row_data):
            try:
                pdf.cell(w, 7, val[:50] if isinstance(val, str) else val, border=1, align='C', fill=True)
            except UnicodeEncodeError:
                # Fallback if fpdf2 cannot encode emojis (latin-1 fallback)
                val_fallback = val.replace('🚶', 'Org:').replace('🪑', 'Bkg:').replace('🗑️', 'Smp:').replace('🌳', 'Phn:').replace('🚦', 'Lmp:')
                pdf.cell(w, 7, val_fallback[:50] if isinstance(val_fallback, str) else val_fallback, border=1, align='C', fill=True)
        
        # Color specific cell for status
        c_r, c_g, c_b = STATUS_COLORS_RGB.get(audit.status, (255, 255, 255))
        pdf.set_fill_color(c_r, c_g, c_b)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font('Helvetica', 'B', 8)
        pdf.cell(col_widths[-1], 7, audit.status, border=1, align='C', fill=True)
        
        pdf.ln()

    output = io.BytesIO(pdf.output())
    output.seek(0)
    return send_file(output, mimetype='application/pdf',
                     as_attachment=True, download_name='Laporan_Audit.pdf')


@main.route('/download/map')
@login_required
def download_map():
    """Generate peta HTML Leaflet sesuai filter yang aktif."""
    audits, _, _ = get_filtered_audits(current_user.id)

    STATUS_COLOR = {
        'Pemeliharaan': '#2ecc71',
        'Monitoring':   '#f1c40f',
        'Perencanaan':  '#ecf0f1',
        'Intervensi':   '#e67e22',
    }

    markers_js = []
    for a in audits:
        if not a.lat or not a.long:
            continue
        color   = STATUS_COLOR.get(a.status, '#e74c3c')
        opacity = max(0.4, min(1.0, a.skor / 100.0))
        total   = ((a.d_orang or 0) + (a.d_bangku or 0) + (a.d_sampah or 0) + (a.d_pohon or 0) + (a.d_lampu or 0) +
                   (a.k_orang or 0) + (a.k_bangku or 0) + (a.k_sampah or 0) + (a.k_pohon or 0) + (a.k_lampu or 0) +
                   (a.ki_orang or 0) + (a.ki_bangku or 0) + (a.ki_sampah or 0) + (a.ki_pohon or 0) + (a.ki_lampu or 0))
        g_street = f"https://www.google.com/maps?q&layer=c&cbll={a.lat},{a.long}"
        g_nav    = f"https://www.google.com/maps/dir/?api=1&destination={a.lat},{a.long}"
        popup_html = (
            f"<b style='font-size:13px'>{a.lokasi}</b><br>"
            f"<span style='color:#888;font-size:11px'>{a.tanggal.strftime('%d/%m/%Y')}</span><br><hr style='margin:4px 0'>"
            f"Skor: <b style='color:{color}'>{a.skor}</b> &nbsp;|&nbsp; Status: <b>{a.status}</b><br>"
            f"Orang: {(a.d_orang or 0)+(a.k_orang or 0)+(a.ki_orang or 0)} &nbsp;"
            f"Bangku: {(a.d_bangku or 0)+(a.k_bangku or 0)+(a.ki_bangku or 0)} &nbsp;"
            f"Sampah: {(a.d_sampah or 0)+(a.k_sampah or 0)+(a.ki_sampah or 0)} &nbsp;"
            f"Pohon: {(a.d_pohon or 0)+(a.k_pohon or 0)+(a.ki_pohon or 0)} &nbsp;"
            f"Lampu: {(a.d_lampu or 0)+(a.k_lampu or 0)+(a.ki_lampu or 0)}<br>"
            f"Total Terdeteksi: <b>{total}</b><br><br>"
            f"<a href='{g_street}' target='_blank'>🔭 Street View</a> &nbsp;&nbsp; "
            f"<a href='{g_nav}' target='_blank'>🧭 Navigasi</a>"
        )
        marker_js = (
            f"L.circleMarker([{a.lat}, {a.long}], {{"
            f"radius: 12, color: '{color}', fillColor: '{color}', "
            f"fillOpacity: {opacity:.2f}, weight: 2, opacity: 1"
            f"}}).addTo(map).bindPopup(`{popup_html}`);"
        )
        markers_js.append(marker_js)

    html_content = f"""<!DOCTYPE html>
<html><head>
<meta charset="UTF-8">
<title>Peta Sebaran Audit - Smart Audit Cerdas</title>
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css"/>
<style>body,html{{margin:0;padding:0;height:100%}} #map{{height:100vh}}</style>
</head><body>
<div id="map"></div>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<script>
var map = L.map('map').setView([-3.97, 122.51], 13);
L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png',{{
 attribution:'&copy; OpenStreetMap contributors'}}).addTo(map);
{chr(10).join(markers_js)}
</script>
</body></html>"""

    output = io.BytesIO(html_content.encode('utf-8'))
    output.seek(0)
    return send_file(output, mimetype='text/html',
                     as_attachment=True,
                     download_name=f'Peta_Audit_{datetime.now().strftime("%Y%m%d")}.html')


# ─── Backup Data ──────────────────────────────────────────────────────────────

@main.route('/backup', methods=['GET', 'POST'])
@login_required
def backup():
    # Backup page removed — redirect to history
    return redirect(url_for('main.history'))


@main.route('/download/backup-json')
@login_required
def download_backup_json():
    """Download semua data audit sebagai file JSON backup."""
    audits = Audit.query.filter_by(user_id=current_user.id).order_by(Audit.tanggal.desc()).all()
    data = []
    for a in audits:
        data.append({
            'id':         a.id,
            'lokasi':     a.lokasi,
            'lat':        a.lat,
            'long':       a.long,
            'tanggal':    a.tanggal.strftime('%d/%m/%Y %H:%M'),
            'skor':       a.skor,
            'status':     a.status,
            'img_depan':  a.img_depan or '',
            'img_kanan':  a.img_kanan or '',
            'img_kiri':   a.img_kiri  or '',
            'counts': {
                'depan': {'orang': a.d_orang or 0, 'bangku': a.d_bangku or 0, 'sampah': a.d_sampah or 0, 'pohon': a.d_pohon or 0, 'lampu': a.d_lampu or 0},
                'kanan': {'orang': a.k_orang or 0, 'bangku': a.k_bangku or 0, 'sampah': a.k_sampah or 0, 'pohon': a.k_pohon or 0, 'lampu': a.k_lampu or 0},
                'kiri':  {'orang': a.ki_orang or 0, 'bangku': a.ki_bangku or 0, 'sampah': a.ki_sampah or 0, 'pohon': a.ki_pohon or 0, 'lampu': a.ki_lampu or 0},
            }
        })
    backup_json = json.dumps({
        'exported_by': current_user.username,
        'exported_at': datetime.now().strftime('%d/%m/%Y %H:%M'),
        'total_records': len(data),
        'audits': data
    }, indent=2, ensure_ascii=False)
    output = io.BytesIO(backup_json.encode('utf-8'))
    output.seek(0)
    return send_file(output, mimetype='application/json',
                     as_attachment=True,
                     download_name=f'Backup_Audit_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json')


# ─── API: Auto-Scan Single Frame ─────────────────────────────────────────────

@main.route('/api/scan_frame', methods=['POST'])
@login_required
def api_scan_frame():
    """
    Menerima 1 frame JPEG dari kamera realtime, menjalankan YOLO inferensi,
    dan mengembalikan bounding box + ringkasan objek dalam format JSON ringan.
    """
    import numpy as np
    import cv2
    from utils import model_yolo

    frame_file = request.files.get('frame')
    if not frame_file:
        return {'error': 'Tidak ada frame yang dikirim.'}, 400

    try:
        # Decode bytes ke numpy array
        file_bytes = frame_file.read()
        np_arr = np.frombuffer(file_bytes, np.uint8)
        img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)

        if img is None:
            return {'error': 'Frame tidak dapat dibaca sebagai gambar.'}, 400

        # Simpan sementara ke temp file
        import tempfile
        with tempfile.NamedTemporaryFile(suffix='.jpg', delete=False) as tmp:
            tmp_path = tmp.name
            _, buf = cv2.imencode('.jpg', img)
            tmp.write(buf.tobytes())

        # Jalankan YOLO inferensi
        results = model_yolo(tmp_path, conf=0.50, verbose=False)
        os.unlink(tmp_path)  # hapus temp file

        if not results or len(results) == 0:
            return {'detections': [], 'summary': {}}

        result = results[0]

        # Class ID mapping (sesuai data.yaml training best.pt)
        CLASS_MAP = {
            0: 'bangkutaman',
            1: 'lampulalulintas',
            2: 'orang',
            3: 'pohonpeneduh',
            4: 'tempatsampah'
        }

        detections = []
        summary    = {'orang': 0, 'bangkutaman': 0, 'tempatsampah': 0, 'pohonpeneduh': 0, 'lampulalulintas': 0}

        for box in result.boxes:
            cls_id = int(box.cls)
            cls_name = CLASS_MAP.get(cls_id, 'unknown')
            conf = float(box.conf)
            x1, y1, x2, y2 = [float(v) for v in box.xyxy[0]]

            detections.append({
                'class_name': cls_name,
                'confidence': round(conf, 3),
                'x1': round(x1), 'y1': round(y1),
                'x2': round(x2), 'y2': round(y2),
            })
            if cls_name in summary:
                summary[cls_name] += 1

        # Remove zeros from summary for a clean response
        summary = {k: v for k, v in summary.items() if v > 0}

        return {'detections': detections, 'summary': summary}

    except Exception as e:
        import traceback
        traceback.print_exc()
        return {'error': str(e)}, 500